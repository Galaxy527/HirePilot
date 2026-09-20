"""Resume text extraction + upload content validation. OCR stubbed."""
from __future__ import annotations

import logging
import zipfile
from pathlib import Path

logger = logging.getLogger(__name__)

ALLOWED_EXT = {"pdf", "txt", "md", "docx"}


class OCRProvider:
    """Extension point — Baidu OCR not implemented this phase."""

    def extract(self, file_path: Path) -> str | None:
        raise NotImplementedError("OCR not enabled this phase")


class BaiduOCRStub(OCRProvider):
    def extract(self, file_path: Path) -> str | None:
        logger.info("BaiduOCRStub called for %s — skipped (not implemented)", file_path)
        return None


class UploadValidationError(ValueError):
    pass


def validate_upload(file_storage, *, allowed_ext: set[str] | None = None) -> tuple[str, bytes]:
    """
    Validate uploaded file by extension + magic bytes.
    Returns (secure-ish original filename, raw bytes). Does not save.
    """
    from werkzeug.utils import secure_filename

    allowed = allowed_ext or ALLOWED_EXT
    if not file_storage or not file_storage.filename:
        raise UploadValidationError("未选择文件")
    filename = secure_filename(file_storage.filename)
    if not filename or "." not in filename:
        raise UploadValidationError("文件名无效")
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in allowed:
        raise UploadValidationError("仅支持 PDF / DOCX / TXT / MD")

    data = file_storage.read()
    if hasattr(file_storage, "seek"):
        file_storage.seek(0)
    if not data:
        raise UploadValidationError("文件为空")

    if ext == "pdf":
        if not data.startswith(b"%PDF"):
            raise UploadValidationError("文件内容不是有效 PDF（魔数校验失败）")
    elif ext == "docx":
        if not data.startswith(b"PK"):
            raise UploadValidationError("文件内容不是有效 DOCX（需为 OOXML/ZIP）")
        try:
            with zipfile.ZipFile(__import__("io").BytesIO(data)) as zf:
                names = set(zf.namelist())
            if "word/document.xml" not in names:
                raise UploadValidationError("DOCX 缺少 word/document.xml")
        except UploadValidationError:
            raise
        except Exception as exc:
            raise UploadValidationError("DOCX 无法解析为有效压缩包") from exc
    else:
        # text / markdown: must be decodable as UTF-8 (allow BOM)
        sample = data[:4096]
        try:
            sample.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise UploadValidationError("文本文件须为 UTF-8 编码") from exc
        # reject obvious binary
        if b"\x00" in sample:
            raise UploadValidationError("文本文件包含二进制空字节")

    return filename, data


def extract_text_from_file(file_path: Path) -> str:
    """Parse PDF / DOCX / plain text. OCR reserved via BaiduOCRStub."""
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(f"Unsupported file type: {suffix}")


def extract_text_from_bytes(filename: str, data: bytes) -> str:
    """Parse from in-memory bytes after validation."""
    import tempfile

    ext = filename.rsplit(".", 1)[-1].lower()
    if ext == "pdf":
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            tmp.write(data)
            path = Path(tmp.name)
        try:
            text = _extract_pdf(path)
        finally:
            path.unlink(missing_ok=True)
        if not text.strip():
            raise UploadValidationError("PDF 无可提取文本（扫描版需 OCR，本期未启用）")
        return text
    if ext == "docx":
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            tmp.write(data)
            path = Path(tmp.name)
        try:
            text = _extract_docx(path)
        finally:
            path.unlink(missing_ok=True)
        if not text.strip():
            raise UploadValidationError("DOCX 无可提取文本")
        return text
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UploadValidationError("无法以 UTF-8 解码文本") from exc


def _extract_pdf(file_path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(file_path))
    parts: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text)
    if not parts:
        logger.warning("PDF had no extractable text; OCR stub reserved for %s", file_path)
        BaiduOCRStub().extract(file_path)
        return ""
    return "\n".join(parts)


def _extract_docx(file_path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise UploadValidationError("服务器未安装 python-docx，无法解析 Word") from exc

    doc = Document(str(file_path))
    parts: list[str] = []
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n".join(parts)
